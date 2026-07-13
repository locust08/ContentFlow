from pathlib import Path
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "ContentFlow_AI_FYP_Proposal_Reference_Style.md"
OUT = ROOT / "ContentFlow_AI_FYP_Proposal_QA_Final.docx"
FIG = ROOT / "_proposal_figures"
FIG.mkdir(exist_ok=True)

FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(90, 90, 90)


def pil_font(size=28, bold=False):
    candidates = [
        r"C:\Windows\Fonts\timesbd.ttf" if bold else r"C:\Windows\Fonts\times.ttf",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def wrap(draw, text, font, max_width):
    lines = []
    for para in text.split("\n"):
        line = ""
        for word in para.split():
            test = f"{line} {word}".strip()
            if draw.textbbox((0, 0), test, font=font)[2] <= max_width or not line:
                line = test
            else:
                lines.append(line)
                line = word
        if line:
            lines.append(line)
    return lines


def draw_box(draw, xy, title, body="", fill=(255, 255, 255)):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=12, fill=fill, outline=(60, 60, 60), width=3)
    draw.text((x1 + 16, y1 + 12), title, font=pil_font(28, True), fill=(0, 0, 0))
    y = y1 + 50
    for line in wrap(draw, body, pil_font(21), x2 - x1 - 32):
        draw.text((x1 + 16, y), line, font=pil_font(21), fill=(35, 35, 35))
        y += 27


def arrow(draw, start, end):
    draw.line([start, end], fill=(40, 40, 40), width=4)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) >= abs(ey - sy):
        sign = 1 if ex > sx else -1
        pts = [(ex, ey), (ex - 18 * sign, ey - 10), (ex - 18 * sign, ey + 10)]
    else:
        sign = 1 if ey > sy else -1
        pts = [(ex, ey), (ex - 10, ey - 18 * sign), (ex + 10, ey - 18 * sign)]
    draw.polygon(pts, fill=(40, 40, 40))


def make_figures():
    def label(draw, xy, text, size=22, bold=False, fill=(70, 70, 70), anchor=None):
        draw.text(xy, text, font=pil_font(size, bold), fill=fill, anchor=anchor)

    def line_label(draw, start, end, text=None):
        arrow(draw, start, end)
        if text:
            mx = (start[0] + end[0]) // 2
            my = (start[1] + end[1]) // 2
            w = draw.textbbox((0, 0), text, font=pil_font(18, True))[2]
            draw.rounded_rectangle((mx - w // 2 - 10, my - 18, mx + w // 2 + 10, my + 12), radius=8, fill=(255, 255, 255), outline=(220, 220, 220))
            label(draw, (mx, my - 15), text, 18, True, anchor="ma")

    # Figure 1: cleaner left-to-right layered architecture.
    img = Image.new("RGB", (1900, 1100), "white")
    d = ImageDraw.Draw(img)
    label(d, (950, 45), "ContentFlow AI System Architecture", 44, True, (0, 0, 0), "ma")
    columns = [
        (70, 120, 380, 980, "1. Interfaces"),
        (440, 120, 750, 980, "2. API & Workflow"),
        (810, 120, 1120, 980, "3. DB & Storage"),
        (1180, 120, 1490, 980, "4. Processing"),
        (1550, 120, 1860, 980, "5. Outputs"),
    ]
    for x1, y1, x2, y2, title in columns:
        d.rounded_rectangle((x1, y1, x2, y2), radius=18, fill=(250, 250, 250), outline=(220, 220, 220), width=2)
        label(d, ((x1 + x2) // 2, y1 + 24), title, 24, True, anchor="ma")
    draw_box(d, (105, 230, 345, 390), "Web Admin", "Manage clients, campaigns, projects, approvals and analytics", (255, 255, 255))
    draw_box(d, (105, 550, 345, 710), "Mobile Module", "Upload assets, view assigned projects and submit review feedback", (255, 255, 255))
    draw_box(d, (475, 390, 715, 550), "Node.js API", "Validation, workflow routing, uploads and service calls", (255, 255, 255))
    draw_box(d, (845, 230, 1085, 390), "Supabase DB", "Users, clients, campaigns, projects, approvals and events", (255, 255, 255))
    draw_box(d, (845, 550, 1085, 710), "Media Storage", "Reference videos, product assets, reaction assets and MP4 files", (255, 255, 255))
    draw_box(d, (1215, 180, 1455, 330), "AI Services", "Transcription, analysis, ideas, scripts and highlights", (255, 255, 255))
    draw_box(d, (1215, 410, 1455, 560), "FFmpeg", "Audio extraction, metadata and media conversion", (255, 255, 255))
    draw_box(d, (1215, 640, 1455, 790), "Remotion", "Programmatic vertical MP4 rendering", (255, 255, 255))
    draw_box(d, (1585, 310, 1825, 470), "Analytics View", "Project totals, renders, approvals and monthly trends", (255, 255, 255))
    draw_box(d, (1585, 610, 1825, 770), "Final Outputs", "Approved clips, render library and delivery-ready content", (255, 255, 255))
    line_label(d, (345, 310), (475, 450), "request")
    line_label(d, (345, 630), (475, 500), "upload")
    line_label(d, (715, 430), (845, 310), "records")
    line_label(d, (715, 510), (845, 630), "files")
    line_label(d, (1085, 310), (1215, 255), "data")
    line_label(d, (1085, 630), (1215, 485), "media")
    line_label(d, (1455, 255), (1585, 390), "metrics")
    line_label(d, (1335, 560), (1335, 640), "processed media")
    line_label(d, (1455, 715), (1585, 690), "render")
    img.save(FIG / "architecture.png")

    # Figure 2: branch-and-merge flow with a single reading direction.
    img = Image.new("RGB", (1900, 1250), "white")
    d = ImageDraw.Draw(img)
    label(d, (950, 45), "ContentFlow AI System Flow", 44, True, (0, 0, 0), "ma")
    top = [
        ((110, 150, 390, 300), "1. Create Client", "Manager records client profile"),
        ((515, 150, 795, 300), "2. Create Campaign", "Campaign objective and brief are defined"),
        ((920, 150, 1200, 300), "3. Create Project", "Type, staff, priority and status are stored"),
        ((1325, 150, 1605, 300), "4. Upload Assets", "Staff uploads references, product files and reaction assets"),
    ]
    for xy, title, body in top:
        draw_box(d, xy, title, body, (255, 255, 255))
    for a, b in [((390, 225), (515, 225)), ((795, 225), (920, 225)), ((1200, 225), (1325, 225))]:
        line_label(d, a, b)
    draw_box(d, (760, 410, 1140, 560), "5. Select Workflow", "Editor chooses AI Generator or Auto Clipper based on project type", (245, 245, 245))
    line_label(d, (1465, 300), (950, 410), "project assets ready")
    draw_box(d, (320, 700, 720, 890), "6A. AI Generator Path", "Analyze reference video, generate content ideas, scripts, prompts and edit plan", (255, 255, 255))
    draw_box(d, (1180, 700, 1580, 890), "6B. Auto Clipper Path", "Transcribe source video, detect highlights, select candidate and subtitle plan", (255, 255, 255))
    line_label(d, (860, 560), (520, 700), "AI Generator")
    line_label(d, (1040, 560), (1380, 700), "Auto Clipper")
    draw_box(d, (760, 980, 1140, 1120), "7. Render, Review & Report", "Render MP4, reviewer approves or requests changes, manager views analytics", (245, 245, 245))
    line_label(d, (520, 890), (820, 980), "edit plan")
    line_label(d, (1380, 890), (1080, 980), "clip plan")
    img.save(FIG / "system_flow.png")

    # Figure 3: cleaner hub-and-spoke ERD with relationship labels.
    img = Image.new("RGB", (1900, 1500), "white")
    d = ImageDraw.Draw(img)
    label(d, (950, 45), "ContentFlow AI Entity Relationship Diagram", 42, True, (0, 0, 0), "ma")
    entities = {
        "users": ("User", ["user_id PK", "org_id FK", "name", "email", "role"]),
        "clients": ("Client", ["client_id PK", "org_id FK", "name", "industry", "contact"]),
        "campaigns": ("Campaign", ["campaign_id PK", "client_id FK", "name", "objective", "status"]),
        "projects": ("Project", ["project_id PK", "campaign_id FK", "assigned_user_id FK", "type", "status", "priority"]),
        "assets": ("Asset", ["asset_id PK", "project_id FK", "asset_type", "file_url", "metadata"]),
        "clip_candidates": ("Clip Candidate", ["candidate_id PK", "project_id FK", "start_time", "end_time", "score", "reason"]),
        "render_jobs": ("Render Job", ["render_id PK", "project_id FK", "output_url", "status", "rendered_at"]),
        "approvals": ("Approval", ["approval_id PK", "project_id FK", "reviewer_id FK", "status", "feedback"]),
        "analytics_events": ("Analytics Event", ["event_id PK", "project_id FK", "user_id FK", "event_type", "created_at"]),
    }
    pos = {
        "clients": (90, 160, 430, 430),
        "campaigns": (560, 160, 900, 430),
        "users": (90, 560, 430, 830),
        "projects": (710, 520, 1080, 830),
        "assets": (1240, 130, 1620, 400),
        "clip_candidates": (1240, 455, 1620, 755),
        "render_jobs": (1240, 810, 1620, 1110),
        "approvals": (560, 1010, 940, 1300),
        "analytics_events": (1030, 1010, 1410, 1300),
    }
    for key, (title, fields) in entities.items():
        x1, y1, x2, y2 = pos[key]
        d.rectangle((x1, y1, x2, y2), fill=(255, 255, 255), outline=(80, 80, 80), width=3)
        d.rectangle((x1, y1, x2, y1 + 48), fill=(242, 242, 242), outline=(80, 80, 80), width=2)
        d.text((x1 + 14, y1 + 10), title, font=pil_font(25, True), fill=(0, 0, 0))
        y = y1 + 62
        for field in fields:
            d.text((x1 + 18, y), field, font=pil_font(21), fill=(25, 25, 25))
            y += 34
    def center(key, side):
        x1, y1, x2, y2 = pos[key]
        return {"l": (x1, (y1 + y2) // 2), "r": (x2, (y1 + y2) // 2), "t": ((x1 + x2) // 2, y1), "b": ((x1 + x2) // 2, y2)}[side]

    def elbow(points, text, label_at):
        for start, end in zip(points[:-2], points[1:-1]):
            d.line([start, end], fill=(40, 40, 40), width=4)
        arrow(d, points[-2], points[-1])
        lx, ly = label_at
        w = d.textbbox((0, 0), text, font=pil_font(18, True))[2]
        d.rounded_rectangle((lx - w // 2 - 10, ly - 18, lx + w // 2 + 10, ly + 12), radius=8, fill=(255, 255, 255), outline=(220, 220, 220))
        label(d, (lx, ly - 15), text, 18, True, anchor="ma")

    elbow([center("clients", "r"), center("campaigns", "l")], "1:M", (500, 292))
    elbow([center("campaigns", "b"), (730, 475), (895, 475), center("projects", "t")], "1:M", (805, 462))
    elbow([center("users", "r"), (570, 695), (570, 675), center("projects", "l")], "1:M assigned", (575, 660))
    elbow([(1080, 570), (1145, 570), (1145, 265), center("assets", "l")], "1:M", (1135, 435))
    elbow([(1080, 650), center("clip_candidates", "l")], "1:M", (1155, 620))
    elbow([(1080, 750), (1145, 750), (1145, 960), center("render_jobs", "l")], "1:M", (1135, 845))
    elbow([(820, 830), (820, 920), (750, 920), center("approvals", "t")], "1:M", (800, 920))
    elbow([(990, 830), (990, 920), (1220, 920), center("analytics_events", "t")], "1:M", (1090, 920))
    elbow([center("users", "b"), (260, 1120), center("approvals", "l")], "1:M reviewer", (350, 1105))
    d.rounded_rectangle((90, 1340, 1810, 1445), radius=14, fill=(250, 250, 250), outline=(200, 200, 200), width=2)
    label(d, (115, 1360), "Cardinality Legend", 22, True)
    label(d, (115, 1395), "1:1 = One-to-One   |   1:M = One-to-Many   |   M:N = Many-to-Many", 21)
    label(d, (980, 1378), "Note: Current prototype uses 1:M relationships.", 19)
    label(d, (980, 1410), "If future M:N staff-project assignment is required, add a project_members junction table.", 19)
    img.save(FIG / "erd.png")


def set_run(run, size=12, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def spacing(p, before=0, after=6, line=1.5):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def add_toc_field(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    for kind, text in [("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None), (None, "Table of contents will update in Microsoft Word."), ("end", None)]:
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText" if text.startswith("TOC") else "w:t")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def add_page_number_footer(section, number_format="decimal", start=1):
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = str(start)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in [begin, instr, separate, text, end]:
        run._r.append(element)
    set_run(run, size=10)

    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))
    pg_num_type.set(qn("w:fmt"), number_format)


def table_borders(table, color="808080"):
    borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table._tbl.tblPr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def shade(cell, fill="F2F2F2"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_text(cell, text, bold=False, size=8, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    spacing(p, after=0, line=1.0)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold)


def gantt(doc):
    weeks = [str(i) for i in range(1, 19)]
    tasks = [
        "Problem identification and proposal planning",
        "Requirement analysis and target user study",
        "System architecture and database design",
        "Web dashboard interface design",
        "Mobile module interface design",
        "Supabase database setup and API integration",
        "Client, campaign, project, and asset management",
        "AI Generator module development",
        "Auto Clipper module development",
        "Mobile upload and review workflow",
        "Rendering and subtitle workflow",
        "Analytics dashboard and reporting",
        "Functional, database, and mobile testing",
        "User evaluation and improvement",
        "Final report, documentation, and presentation",
    ]
    ranges = [(1, 5), (2, 5), (3, 6), (4, 7), (4, 7), (6, 9), (7, 10), (8, 11), (8, 12), (9, 12), (10, 13), (11, 14), (12, 15), (13, 16), (14, 18)]
    table = doc.add_table(rows=len(tasks) + 1, cols=19)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(table)
    cell_text(table.cell(0, 0), "Task", True, 8, WD_ALIGN_PARAGRAPH.LEFT)
    shade(table.cell(0, 0))
    for i, week in enumerate(weeks, 1):
        cell_text(table.cell(0, i), week, True, 8)
        shade(table.cell(0, i))
    for r, task in enumerate(tasks, 1):
        cell_text(table.cell(r, 0), task, False, 7, WD_ALIGN_PARAGRAPH.LEFT)
        a, b = ranges[r - 1]
        for c in range(1, 19):
            cell_text(table.cell(r, c), "X" if a <= c <= b else "", False, 7)


def erd_relationship_table(doc):
    rows = [
        ("Client", "Campaign", "1:M", "One client can have many campaigns; each campaign belongs to one client."),
        ("Campaign", "Project", "1:M", "One campaign can contain many content projects; each project belongs to one campaign."),
        ("User", "Project", "1:M", "One user can be assigned to many projects; each project stores one main assigned user in the prototype."),
        ("Project", "Asset", "1:M", "One project can contain many uploaded assets such as product images, references, and reaction files."),
        ("Project", "Clip Candidate", "1:M", "One project can generate many highlight or clip candidates."),
        ("Project", "Render Job", "1:M", "One project can have many render jobs or rendered MP4 outputs."),
        ("Project", "Approval", "1:M", "One project can have many approval or review records."),
        ("User", "Approval", "1:M", "One reviewer can submit many approval records; each approval is linked to one reviewer."),
        ("Project", "Analytics Event", "1:M", "One project can generate many analytics events for reporting and monitoring."),
        ("User", "Project via project_members", "M:N (future optional)", "If multiple staff members need to work on the same project, a junction table can connect many users to many projects."),
    ]
    table = doc.add_table(rows=len(rows) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(table)
    headers = ["Parent Entity", "Child Entity", "Relation", "Explanation"]
    widths = [1.2, 1.35, 0.85, 3.1]
    for c, header in enumerate(headers):
        cell_text(table.cell(0, c), header, True, 8, WD_ALIGN_PARAGRAPH.CENTER)
        shade(table.cell(0, c))
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if c == 2 else WD_ALIGN_PARAGRAPH.LEFT
            cell_text(table.cell(r, c), value, False, 7.5, align)
            table.cell(r, c).width = Inches(widths[c])


def add_para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, bold=False, italic=False, size=12):
    p = doc.add_paragraph()
    p.alignment = align
    spacing(p, after=after)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, italic=italic)


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    spacing(p, before=12 if level == 1 else 8, after=6 if level < 3 else 3, line=1.15)
    r = p.add_run(text.upper() if level < 3 else text)
    set_run(r, bold=True)


def build():
    make_figures()
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    add_page_number_footer(doc.sections[0], "lowerRoman", 1)
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"]:
        st = doc.styles[style_name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        st.font.size = Pt(12)
        st.font.color.rgb = BLACK
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.bold = True
    doc.styles["Heading 3"].font.bold = True

    lines = SRC.read_text(encoding="utf-8").splitlines()
    cover = []
    while lines:
        line = lines.pop(0)
        if line.strip() == r"\pagebreak":
            break
        if line.strip():
            cover.append(re.sub(r"^#+\s*", "", line.strip()))
    add_para(doc, "", after=12)
    for line in cover:
        bold = line.isupper() or line.endswith(":") or line.startswith("CONTENTFLOW")
        add_para(doc, line, WD_ALIGN_PARAGRAPH.CENTER, after=8 if line.endswith(":") else 12, bold=bold)
    doc.add_page_break()

    pending_break = False
    main_started = False
    for raw in lines:
        s = raw.strip()
        if not s:
            continue
        if s == r"\pagebreak":
            pending_break = True
            continue
        if pending_break:
            if s.startswith("# 1. ") and not main_started:
                section = doc.add_section(WD_SECTION.NEW_PAGE)
                section.top_margin = Inches(1)
                section.bottom_margin = Inches(1)
                section.left_margin = Inches(1)
                section.right_margin = Inches(1)
                add_page_number_footer(section, "decimal", 1)
                main_started = True
            else:
                doc.add_page_break()
            pending_break = False
        if s == "[AUTO_TOC]":
            add_toc_field(doc)
        elif s == "[GANTT_TABLE]":
            gantt(doc)
        elif s == "[ERD_RELATIONSHIP_TABLE]":
            erd_relationship_table(doc)
        elif s.startswith("[FIGURE:"):
            name = s[8:-1]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(FIG / f"{name}.png"), width=Inches(6.2))
        elif s.startswith("# "):
            add_heading(doc, s[2:].strip(), 1)
        elif s.startswith("## "):
            add_heading(doc, s[3:].strip(), 2)
        elif s.startswith("### "):
            add_heading(doc, s[4:].strip(), 3)
        elif re.match(r"^\d+\.\s+", s):
            p = doc.add_paragraph(style="List Number")
            spacing(p, after=3)
            r = p.add_run(re.sub(r"^\d+\.\s+", "", s))
            set_run(r)
        else:
            italic = s.startswith("Figure ") or s.startswith("Table ")
            align = WD_ALIGN_PARAGRAPH.CENTER if italic else WD_ALIGN_PARAGRAPH.JUSTIFY
            add_para(doc, s, align=align, after=6, italic=italic, size=11 if italic else 12)
    doc.core_properties.title = "ContentFlow AI FYP Proposal"
    doc.core_properties.author = "Codex"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
